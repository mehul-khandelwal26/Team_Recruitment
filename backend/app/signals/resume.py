"""Resume Signal Extractor for PDF and DOCX candidate resumes.

Extracts evidence-backed information from resumes without OCR:
- Skills and Technologies (programming languages, frameworks, databases, tools)
- Projects (titles, descriptions, associated technologies, and evidence snippets)
- Education (degrees, institutions, graduation years, GPA)
- Experience and Achievements (internships, work history, roles, certifications)
- External Links (GitHub, portfolio, LinkedIn, other professional profiles)

Includes SSRF-safe URL downloading, document validation, error isolation,
and integration with BaseSignalExtractor and SignalExtractorRegistry.
"""

from __future__ import annotations

import datetime
import io
import ipaddress
import os
import re
import socket
import urllib.parse
from typing import Any, Optional

import docx
import httpcore
import httpx
from pypdf import PdfReader
from pypdf.errors import PdfReadError

from app.schemas.ingestion import MemberSchema
from app.signals.base import BaseSignalExtractor
from app.signals.schemas import (
    ExtractedSignal,
    ExtractorDescriptor,
    SignalDefinition,
    SignalEvidence,
    SignalStatus,
)

# ---------------------------------------------------------------------------
# Maximum sizes and timeouts
# ---------------------------------------------------------------------------

DEFAULT_MAX_RESUME_BYTES = 10 * 1024 * 1024  # 10 MB
DEFAULT_DOWNLOAD_TIMEOUT = 10.0  # seconds

# ---------------------------------------------------------------------------
# Technology and Skill Catalogs
# ---------------------------------------------------------------------------

PROGRAMMING_LANGUAGES: dict[str, str] = {
    r"\bpython\b": "Python",
    r"\bjavascript\b|\bjs\b": "JavaScript",
    r"\btypescript\b|\bts\b": "TypeScript",
    r"\bc\+\+(?!\w)|\bcpp\b": "C++",
    r"\bc#(?!\w)|\bc-sharp\b": "C#",
    r"\b(?:c\s+programming|ansi\s+c|c\s*[/,]\s*c\+\+|c\s*[/,]\s*cpp)\b": "C",
    r"\bjava\b(?!script)": "Java",
    r"\bgolang\b|\bgo\s+(?:programming|developer|engineer|lang)\b|(?:(?<=[,/|•:-])\s*go(?:\s*[,/|•:-]|\b))|(?:\bgo\s*[,/|•:-])": "Go",
    r"\brust\b": "Rust",
    r"\bruby\b": "Ruby",
    r"\bphp\b": "PHP",
    r"\bswift\b": "Swift",
    r"\bkotlin\b": "Kotlin",
    r"\bscala\b": "Scala",
    r"\br\s+programming\b|\br\s*[,/]\s*python\b|\br-lang\b": "R",
    r"\bdart\b": "Dart",
    r"\bsql\b|\bpl/sql\b|\bt-sql\b": "SQL",
    r"\bbash\b|\bshell\b|\bpowershell\b": "Shell",
    r"\bhtml(?:5)?\b": "HTML",
    r"\bcss(?:3)?\b": "CSS",
    r"\bsolidity\b": "Solidity",
    r"\bmatlab\b": "MATLAB",
    r"\bperl\b": "Perl",
    r"\bhaskell\b": "Haskell",
    r"\bjulia\b": "Julia",
    r"\blua\b": "Lua",
}

FRAMEWORKS_LIBRARIES: dict[str, str] = {
    r"\breact(?:\.js|js)?\b(?! native)": "React",
    r"\breact\s+native\b": "React Native",
    r"\bnext(?:\.js|js)?\b": "Next.js",
    r"\bvue(?:\.js|js)?\b": "Vue",
    r"\bangular(?:\.js|js)?\b": "Angular",
    r"\bsvelte\b": "Svelte",
    r"\bdjango\b": "Django",
    r"\bfastapi\b": "FastAPI",
    r"\bflask\b": "Flask",
    r"\bspring\s+boot\b|\bspring\s+framework\b": "Spring Boot",
    r"\bexpress(?:\.js|js)?\b": "Express",
    r"\bnode(?:\.js|js)?\b": "Node.js",
    r"\bnest(?:\.js|js)?\b": "NestJS",
    r"\basp\.net\b|\b\.net\s+core\b|\b\.net\b": ".NET",
    r"\bruby\s+on\s+rails\b|\brails\b": "Ruby on Rails",
    r"\bpytorch\b": "PyTorch",
    r"\btensorflow\b": "TensorFlow",
    r"\bkeras\b": "Keras",
    r"\bscikit-learn\b|\bsklearn\b": "Scikit-learn",
    r"\bpandas\b": "Pandas",
    r"\bnumpy\b": "NumPy",
    r"\btailwind(?:\s+css)?\b": "Tailwind CSS",
    r"\bbootstrap\b": "Bootstrap",
    r"\bgraphql\b": "GraphQL",
    r"\bredux\b": "Redux",
    r"\bopencv\b": "OpenCV",
}

DATABASES: dict[str, str] = {
    r"\bpostgresql\b|\bpostgres\b": "PostgreSQL",
    r"\bmysql\b": "MySQL",
    r"\bsqlite(?:3)?\b": "SQLite",
    r"\bmongodb\b|\bmongo\b": "MongoDB",
    r"\bredis\b": "Redis",
    r"\bcassandra\b": "Cassandra",
    r"\bdynamodb\b": "DynamoDB",
    r"\boracle\s+db\b|\boracle\s+database\b": "Oracle",
    r"\belasticsearch\b": "Elasticsearch",
    r"\bfirebase\b": "Firebase",
    r"\bsupabase\b": "Supabase",
    r"\bsnowflake\b": "Snowflake",
    r"\bbigquery\b": "BigQuery",
    r"\bmariadb\b": "MariaDB",
    r"\bneo4j\b": "Neo4j",
}

TOOLS_TECHNOLOGIES: dict[str, str] = {
    r"\bdocker\b": "Docker",
    r"\bkubernetes\b|\bk8s\b": "Kubernetes",
    r"\bgit\b(?!hub|lab)": "Git",
    r"\bgithub\b": "GitHub",
    r"\bgitlab\b": "GitLab",
    r"\blinux\b|\bubuntu\b|\bdebian\b": "Linux",
    r"\baws\b|\bamazon\s+web\s+services\b": "AWS",
    r"\bazure\b": "Azure",
    r"\bgcp\b|\bgoogle\s+cloud\b": "GCP",
    r"\bjenkins\b": "Jenkins",
    r"\bci/cd\b|\bcontinuous\s+integration\b": "CI/CD",
    r"\bterraform\b": "Terraform",
    r"\bansible\b": "Ansible",
    r"\bnginx\b": "Nginx",
    r"\bapache\b": "Apache",
    r"\bpostman\b": "Postman",
    r"\bkafka\b": "Kafka",
    r"\brabbitmq\b": "RabbitMQ",
    r"\bjira\b": "Jira",
    r"\bfigma\b": "Figma",
}

# ---------------------------------------------------------------------------
# SSRF-Safe Downloader and DNS Rebinding Protection
# ---------------------------------------------------------------------------


class SSRFSecurityError(Exception):
    """Raised when an outbound network connection violates SSRF security policy."""
    pass


def _is_unc_path(path: str) -> bool:
    r"""Detect UNC network share paths (e.g. \\server\share or //server/share or file://server/share)."""
    s = path.strip()
    if s.startswith(("\\\\", "//")):
        return True
    if s.lower().startswith("file://"):
        try:
            parsed = urllib.parse.urlparse(s)
            if parsed.netloc and parsed.netloc.lower() not in {"", "localhost"}:
                return True
        except Exception:
            return True
    return False


class SSRFSafeNetworkBackend(httpcore.SyncBackend):
    """Network backend enforcing strict port allowlists, IP restrictions, and DNS rebinding defense."""

    def connect_tcp(self, host: str, port: int, **kwargs):
        if port not in {80, 443}:
            raise SSRFSecurityError(
                f"Restricted port '{port}' not allowed. Only standard web ports (80, 443) are allowed."
            )

        # Resolve all IPs before connect
        try:
            addr_infos = socket.getaddrinfo(host, None)
        except Exception as exc:
            raise SSRFSecurityError(f"Failed to resolve hostname '{host}': {exc}")

        if not addr_infos:
            raise SSRFSecurityError(f"No IP addresses could be resolved for host '{host}'.")

        for family, _, _, _, sockaddr in addr_infos:
            ip_str = sockaddr[0]
            try:
                ip_obj = ipaddress.ip_address(ip_str)
            except ValueError:
                raise SSRFSecurityError(f"Invalid resolved IP address '{ip_str}'.")

            if not ip_obj.is_global or str(ip_obj) == "169.254.169.254":
                raise SSRFSecurityError(f"Access to non-global IP address '{ip_str}' is blocked for security.")

        stream = super().connect_tcp(host, port, **kwargs)
        peer = stream.get_extra_info("peername")
        if peer:
            peer_ip_str = peer[0]
            try:
                peer_ip_obj = ipaddress.ip_address(peer_ip_str)
                if not peer_ip_obj.is_global or str(peer_ip_obj) == "169.254.169.254":
                    stream.close()
                    raise SSRFSecurityError(
                        f"Connected peer IP address '{peer_ip_str}' is blocked (DNS rebinding protection)."
                    )
            except ValueError:
                stream.close()
                raise SSRFSecurityError(f"Invalid connected peer IP address '{peer_ip_str}'.")

        return stream


class SSRFSafeTransport(httpx.HTTPTransport):
    """HTTP transport configured with SSRFSafeNetworkBackend to prevent SSRF and DNS rebinding."""

    def __init__(self, **kwargs):
        super().__init__(**kwargs)
        self._pool = httpcore.ConnectionPool(
            network_backend=SSRFSafeNetworkBackend(),
            retries=0,
        )


def validate_url_security(url: str) -> tuple[bool, Optional[str]]:
    """Validate that a URL is safe against SSRF attacks before downloading.

    Enforces:
    - Scheme must be http or https
    - Hostname must be present and resolve to a public IP
    - Private networks (RFC 1918), loopback, link-local, Carrier-Grade NAT (CGNAT),
      and cloud metadata IPs are strictly blocked (verifies ip.is_global)
    - Credentials in URL blocked
    - Non-standard ports restricted strictly to standard web ports {80, 443}
    - Checks ALL resolved IP addresses
    """
    if not url or not isinstance(url, str):
        return False, "URL is empty or not a string."

    try:
        parsed = urllib.parse.urlparse(url.strip())
    except Exception as exc:
        return False, f"Malformed URL: {exc}"

    if parsed.scheme.lower() not in {"http", "https"}:
        return False, f"Invalid URL scheme '{parsed.scheme}'. Only HTTP and HTTPS are permitted."

    if not parsed.netloc:
        return False, "URL lacks a valid network host."

    if parsed.username or parsed.password:
        return False, "URLs with embedded credentials (user:password@) are not permitted."

    # Validate port if specified - strictly 80 and 443
    if parsed.port and parsed.port not in {80, 443}:
        return False, f"Restricted port '{parsed.port}' not allowed. Only standard web ports (80, 443) are allowed."

    hostname = parsed.hostname
    if not hostname:
        return False, "URL lacks a valid hostname."

    # Resolve IP addresses and verify against private/local networks
    try:
        addr_infos = socket.getaddrinfo(hostname, None)
    except socket.gaierror as exc:
        return False, f"Failed to resolve hostname '{hostname}': {exc}"

    if not addr_infos:
        return False, f"No IP addresses could be resolved for host '{hostname}'."

    for family, _, _, _, sockaddr in addr_infos:
        ip_str = sockaddr[0]
        try:
            ip_obj = ipaddress.ip_address(ip_str)
        except ValueError:
            return False, f"Invalid resolved IP address '{ip_str}'."

        if not ip_obj.is_global or str(ip_obj) == "169.254.169.254":
            return False, f"Access to non-global IP address '{ip_str}' is blocked for security."

    return True, None


def download_resume_file(
    url: str,
    max_bytes: int = DEFAULT_MAX_RESUME_BYTES,
    timeout: float = DEFAULT_DOWNLOAD_TIMEOUT,
) -> tuple[Optional[bytes], Optional[str]]:
    """Safely download a resume file over HTTP/HTTPS with size limit, SSRF protection, and rebinding defense."""
    is_safe, err_msg = validate_url_security(url)
    if not is_safe:
        return None, err_msg

    try:
        transport = SSRFSafeTransport()
        with httpx.Client(transport=transport, timeout=timeout, follow_redirects=False) as client:
            with client.stream("GET", url, headers={"User-Agent": "TeamRecruitmentTool-ResumeAnalyzer/1.0"}) as resp:
                if 300 <= resp.status_code < 400:
                    return None, f"HTTP redirects are not permitted for resume downloads (received HTTP {resp.status_code})."

                if resp.status_code >= 400:
                    return None, f"Failed to download resume: HTTP {resp.status_code}"

                content_len = resp.headers.get("content-length")
                if content_len and int(content_len) > max_bytes:
                    return None, f"Resume exceeds maximum allowed size of {max_bytes} bytes."

                buf = io.BytesIO()
                downloaded = 0
                for chunk in resp.iter_bytes(chunk_size=8192):
                    downloaded += len(chunk)
                    if downloaded > max_bytes:
                        return None, f"Resume download exceeded maximum allowed size of {max_bytes} bytes."
                    buf.write(chunk)
                return buf.getvalue(), None
    except SSRFSecurityError as exc:
        return None, f"Security restriction: {exc}"
    except httpx.TimeoutException:
        return None, f"Request timed out while downloading resume after {timeout} seconds."
    except httpx.RequestError as exc:
        if exc.__cause__ and isinstance(exc.__cause__, SSRFSecurityError):
            return None, f"Security restriction: {exc.__cause__}"
        return None, f"Network error while downloading resume: {exc}"
    except Exception as exc:
        return None, f"Unexpected error downloading resume: {exc}"


# ---------------------------------------------------------------------------
# Document Text Extraction (PDF & DOCX)
# ---------------------------------------------------------------------------


def extract_text_from_document(
    raw_bytes: bytes,
    filename_or_url: str = "",
) -> tuple[Optional[str], Optional[str], dict[str, Any]]:
    """Extract plain text from PDF or DOCX file bytes.

    Returns:
        (text, error_message, metadata)
    """
    if not raw_bytes or len(raw_bytes) == 0:
        return None, "Resume document is empty (0 bytes).", {"empty": True}

    # Format detection by magic bytes and extension
    is_pdf = raw_bytes.startswith(b"%PDF-") or filename_or_url.lower().endswith(".pdf")
    is_docx = (raw_bytes.startswith(b"PK\x03\x04") and (filename_or_url.lower().endswith(".docx") or b"word/" in raw_bytes[:2048])) or filename_or_url.lower().endswith(".docx")

    if is_pdf:
        try:
            reader = PdfReader(io.BytesIO(raw_bytes))
            if reader.is_encrypted:
                # Try decrypting with blank password
                try:
                    decrypted = reader.decrypt("")
                except Exception:
                    decrypted = 0
                if decrypted == 0:
                    return (
                        None,
                        "Encrypted PDF document cannot be read without a password.",
                        {"format": "pdf", "encrypted": True},
                    )

            pages_text = []
            for page in reader.pages:
                t = page.extract_text() or ""
                pages_text.append(t)
            full_text = "\n\n".join(pages_text).strip()

            if not full_text or len(re.sub(r"\s+", "", full_text)) < 20:
                return (
                    None,
                    "Document contains no readable text (may be scanned image, empty, or non-text PDF). OCR is not enabled.",
                    {"format": "pdf", "pages": len(reader.pages), "empty_or_scanned": True},
                )

            return full_text, None, {"format": "pdf", "pages": len(reader.pages), "char_count": len(full_text)}
        except PdfReadError as exc:
            return None, f"Corrupted or invalid PDF document: {exc}", {"format": "pdf", "corrupted": True}
        except Exception as exc:
            return None, f"Failed to parse PDF document: {exc}", {"format": "pdf", "error": str(exc)}

    elif is_docx:
        try:
            doc = docx.Document(io.BytesIO(raw_bytes))
            paras = [p.text for p in doc.paragraphs if p.text.strip()]
            table_cells = []
            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        if cell.text.strip():
                            table_cells.append(cell.text.strip())
            full_text = "\n".join(paras + table_cells).strip()

            if not full_text or len(re.sub(r"\s+", "", full_text)) < 20:
                return None, "Document contains no readable text.", {"format": "docx", "empty": True}

            return full_text, None, {"format": "docx", "paragraphs": len(doc.paragraphs), "char_count": len(full_text)}
        except Exception as exc:
            return None, f"Corrupted or invalid DOCX document: {exc}", {"format": "docx", "corrupted": True}

    else:
        return (
            None,
            "Unsupported resume file format. Only PDF (.pdf) and DOCX (.docx) documents are supported.",
            {"unsupported": True},
        )


# ---------------------------------------------------------------------------
# Section and Signal Pattern Extractors
# ---------------------------------------------------------------------------


def _find_snippet_lines(text: str, keywords: list[str], max_lines: int = 4) -> str:
    """Find text lines containing any of the given keywords for evidence snippets."""
    matching: list[str] = []
    lines = [l.strip() for l in text.split("\n") if l.strip()]
    for line in lines:
        for kw in keywords:
            if kw.lower() in line.lower():
                matching.append(line)
                break
        if len(matching) >= max_lines:
            break
    return " | ".join(matching) if matching else ""


AMBIGUOUS_SKILLS: dict[str, str] = {
    "Next.js": r"\bnext(?:\.js|js)\b",
    "Express": r"\bexpress(?:\.js|js)\b",
    "Swift": r"\b(?:swift\s+(?:programming|developer|engineer|ios|lang|code)|apple\s+swift)\b",
    "Rust": r"\b(?:rust\s+(?:programming|developer|engineer|lang|code))\b",
    "Dart": r"\b(?:dart\s+(?:programming|developer|engineer|lang|flutter))\b",
    "Julia": r"\b(?:julia\s+(?:programming|developer|engineer|lang))\b",
    "Ruby": r"\b(?:ruby\s+on\s+rails|ruby\s+(?:programming|developer|lang|gem))\b",
    "Lua": r"\b(?:lua\s+(?:programming|scripting|developer|lang))\b",
    "Shell": r"\b(?:shell\s+script(?:ing)?|powershell|bash\s+shell)\b",
    "Apache": r"\bapache\s+(?:kafka|spark|cassandra|tomcat|hadoop|flink|airflow|server|httpd)\b",
    "Git": r"\bgit\s+(?:version\s+control|workflow|vcs|cli|repo|actions|rebase|merge)\b",
    "Go": r"\b(?:golang|go\s+(?:programming|developer|engineer|lang))\b",
    "Bash": r"\bbash\s+(?:script(?:ing)?|cli|shell)\b",
    "Spring Boot": r"\bspring\s+(?:boot|framework|cloud|data|mvc|security)\b",
    "Flask": r"\bflask\s+(?:framework|app|api|python)\b",
    "Ruby on Rails": r"\b(?:ruby\s+on\s+rails|rails\s+framework|rails\s+app)\b",
}


def _is_tech_list_line(line: str) -> bool:
    """Detect whether a line is a structured technical skill list rather than conversational prose."""
    clean = line.strip()
    if not clean:
        return False
    if re.match(
        r"^(?:skills|technologies|technical\s+skills|tech\s+stack|languages|frameworks|tools|databases|platforms|competencies)\s*[:=-]",
        clean,
        re.IGNORECASE,
    ):
        return True
    parts = [p.strip() for p in re.split(r"[,|•/]", clean) if p.strip()]
    if len(parts) >= 3 and all(
        len(p) <= 30
        and not re.search(r"\b(?:the|and|in|to|my|for|with|of|is|was|i|we|our|you|role|want|aspire|ideas)\b", p.lower())
        for p in parts
    ):
        return True
    return False


def extract_skills_from_text(text: str) -> tuple[dict[str, Any], list[str]]:
    """Extract categorized technical skills and evidence from resume text.

    Ambiguous skill tokens (e.g. next, express, swift, rust, dart, julia, ruby,
    lua, shell, apache, git, go, bash, spring, flask, rails) only match:
    1. Inside an identified Skills or Projects section, OR
    2. When written in canonical form (Next.js, Express.js, Golang, Ruby on Rails, etc.), OR
    3. In a comma-, pipe-, or bullet-separated list of technologies.
    """
    languages: list[str] = []
    frameworks: list[str] = []
    databases: list[str] = []
    tools: list[str] = []

    lines = [l.strip() for l in text.split("\n") if l.strip()]
    in_skills_section = False
    in_projects_section = False

    skills_header_re = re.compile(
        r"^(?:skills|technical\s+skills|core\s+competencies|technologies|key\s+skills)\b", re.IGNORECASE
    )
    project_header_re = re.compile(
        r"^(?:projects?|academic\s+projects?|personal\s+projects?|technical\s+projects?)\b", re.IGNORECASE
    )
    section_break_re = re.compile(
        r"^(?:education|experience|work\s+experience|certifications?|achievements?|activities|summary|objective|profile)\b",
        re.IGNORECASE,
    )

    for line in lines:
        if skills_header_re.search(line):
            in_skills_section = True
            in_projects_section = False
            remainder = re.sub(
                r"^(?:skills|technical\s+skills|core\s+competencies|technologies|key\s+skills)\s*[:=-]?\s*",
                "",
                line,
                flags=re.IGNORECASE,
            ).strip()
            if not remainder:
                continue

        if project_header_re.search(line):
            in_projects_section = True
            in_skills_section = False
            remainder = re.sub(
                r"^(?:projects?|academic\s+projects?|personal\s+projects?|technical\s+projects?)\s*[:=-]?\s*",
                "",
                line,
                flags=re.IGNORECASE,
            ).strip()
            if not remainder:
                continue

        if (in_skills_section or in_projects_section) and section_break_re.search(line):
            in_skills_section = False
            in_projects_section = False
            continue

        allowed_full_match = in_skills_section or in_projects_section or _is_tech_list_line(line)
        line_lower = line.lower()

        if allowed_full_match:
            for pattern, name in PROGRAMMING_LANGUAGES.items():
                if re.search(pattern, line_lower) and name not in languages:
                    languages.append(name)
            for pattern, name in FRAMEWORKS_LIBRARIES.items():
                if re.search(pattern, line_lower) and name not in frameworks:
                    frameworks.append(name)
            for pattern, name in DATABASES.items():
                if re.search(pattern, line_lower) and name not in databases:
                    databases.append(name)
            for pattern, name in TOOLS_TECHNOLOGIES.items():
                if re.search(pattern, line_lower) and name not in tools:
                    tools.append(name)
        else:
            # Free prose line outside skills/projects: match canonical forms for ambiguous skills
            for amb_name, canon_pat in AMBIGUOUS_SKILLS.items():
                if re.search(canon_pat, line_lower):
                    if amb_name in ("Swift", "Rust", "Dart", "Julia", "Ruby", "Lua", "Shell", "Go", "Bash"):
                        if amb_name not in languages:
                            languages.append(amb_name)
                    elif amb_name in ("Next.js", "Express", "Spring Boot", "Flask", "Ruby on Rails"):
                        if amb_name not in frameworks:
                            frameworks.append(amb_name)
                    elif amb_name in ("Apache", "Git"):
                        if amb_name not in tools:
                            tools.append(amb_name)

            # Match unambiguous skills
            for pattern, name in PROGRAMMING_LANGUAGES.items():
                if name not in AMBIGUOUS_SKILLS and name not in languages:
                    if re.search(pattern, line_lower):
                        languages.append(name)
            for pattern, name in FRAMEWORKS_LIBRARIES.items():
                if name not in AMBIGUOUS_SKILLS and name not in frameworks:
                    if re.search(pattern, line_lower):
                        frameworks.append(name)
            for pattern, name in DATABASES.items():
                if name not in databases and re.search(pattern, line_lower):
                    databases.append(name)
            for pattern, name in TOOLS_TECHNOLOGIES.items():
                if name not in AMBIGUOUS_SKILLS and name not in tools:
                    if re.search(pattern, line_lower):
                        tools.append(name)

    all_skills = sorted(list(set(languages + frameworks + databases + tools)))
    evidence_lines = _find_snippet_lines(text, all_skills[:6])

    skill_data = {
        "languages": sorted(languages),
        "frameworks": sorted(frameworks),
        "databases": sorted(databases),
        "tools": sorted(tools),
        "all_skills": all_skills,
        "total_skills_count": len(all_skills),
        "note": "Mention of technologies in resume text does not independently prove proficiency.",
    }
    return skill_data, [evidence_lines] if evidence_lines else []


def extract_projects_from_text(text: str) -> tuple[list[dict[str, Any]], str]:
    """Extract project names, descriptions, associated technologies, and evidence."""
    projects: list[dict[str, Any]] = []

    lines = [l.strip() for l in text.split("\n") if l.strip()]
    in_project_section = False
    current_project: Optional[dict[str, Any]] = None

    project_header_re = re.compile(r"^(?:projects?|academic\s+projects?|personal\s+projects?|technical\s+projects?)\b", re.IGNORECASE)
    section_break_re = re.compile(r"^(?:education|experience|skills|technical\s+skills|certifications?|achievements?|activities)\b", re.IGNORECASE)

    for line in lines:
        if project_header_re.search(line):
            in_project_section = True
            continue
        if in_project_section and section_break_re.search(line):
            in_project_section = False
            if current_project:
                projects.append(current_project)
                current_project = None
            continue

        if in_project_section:
            title_match = re.match(r"^([A-Z0-9][A-Za-z0-9\s_\-:]{3,40})(?:\s*[\|\-\–\—\:]\s*(.+))?$", line)
            if title_match and not line.startswith(("-", "•", "*", "http")):
                if current_project:
                    projects.append(current_project)
                p_name = title_match.group(1).strip()
                p_tech_str = title_match.group(2) or ""

                techs: list[str] = []
                for p_dict in [PROGRAMMING_LANGUAGES, FRAMEWORKS_LIBRARIES, DATABASES, TOOLS_TECHNOLOGIES]:
                    for pat, tname in p_dict.items():
                        if (
                            re.search(pat, p_tech_str.lower())
                            or (tname == "Go" and re.search(r"\bgo\b", p_tech_str.lower()))
                        ) and tname not in techs:
                            techs.append(tname)

                current_project = {
                    "name": p_name,
                    "technologies": techs,
                    "description": "",
                    "evidence": line,
                    "verified": False,
                    "provenance": "self_reported_resume",
                }
            elif current_project:
                desc = line.lstrip("-•* ").strip()
                if not current_project["description"]:
                    current_project["description"] = desc
                else:
                    current_project["description"] += " " + desc
                for p_dict in [PROGRAMMING_LANGUAGES, FRAMEWORKS_LIBRARIES, DATABASES, TOOLS_TECHNOLOGIES]:
                    for pat, tname in p_dict.items():
                        if (
                            re.search(pat, desc.lower())
                            or (tname == "Go" and re.search(r"\b(?:golang|go\s+(?:programming|lang|code))\b", desc.lower()))
                        ) and tname not in current_project["technologies"]:
                            current_project["technologies"].append(tname)

    if current_project:
        projects.append(current_project)

    evidence_summary = " | ".join(p["name"] for p in projects[:3]) if projects else ""
    return projects, evidence_summary


def extract_education_from_text(text: str) -> tuple[dict[str, Any], Optional[str], str]:
    """Extract degree, institution, graduation year, GPA, and evidence snippet."""
    degree: Optional[str] = None
    degree_short: Optional[str] = None
    institution: Optional[str] = None
    grad_year: Optional[int] = None
    gpa: Optional[str] = None
    evidence: str = ""

    lines = [l.strip() for l in text.split("\n") if l.strip()]

    # Strict degree patterns: require dotted abbreviation, exact uppercase BE/BS/MS, or full words
    degree_patterns: list[tuple[str, str, str, int]] = [
        (r"\b(?:b\.?tech(?:nology)?|bachelor\s+of\s+technology)\b", "Bachelor of Technology", "B.Tech", re.IGNORECASE),
        (r"\b(?:b\.e\.|b\.eng\.?|bachelor\s+of\s+engineering)\b", "Bachelor of Engineering", "B.E.", re.IGNORECASE),
        (r"\bBE\b", "Bachelor of Engineering", "B.E.", 0),
        (r"\b(?:b\.s\.|b\.sc\.?|bachelor\s+of\s+science)\b", "Bachelor of Science", "B.S.", re.IGNORECASE),
        (r"\bBS\b", "Bachelor of Science", "B.S.", 0),
        (r"\b(?:bca|bachelor\s+of\s+computer\s+applications)\b", "Bachelor of Computer Applications", "BCA", re.IGNORECASE),
        (r"\b(?:m\.?tech(?:nology)?|master\s+of\s+technology)\b", "Master of Technology", "M.Tech", re.IGNORECASE),
        (r"\b(?:m\.s\.|m\.sc\.?|master\s+of\s+science)\b", "Master of Science", "M.S.", re.IGNORECASE),
        (r"\bMS\b(?!\s*(?:Office|Word|Excel|PowerPoint|Paint|Access|Outlook|SQL|Visual|Project|Teams|DOS|Windows|\.NET))\b", "Master of Science", "M.S.", 0),
        (r"\b(?:mca|master\s+of\s+computer\s+applications)\b", "Master of Computer Applications", "MCA", re.IGNORECASE),
        (r"\b(?:mba|master\s+of\s+business\s+administration)\b", "Master of Business Administration", "MBA", re.IGNORECASE),
        (r"\b(?:ph\.?d\.?|doctor\s+of\s+philosophy)\b", "Doctor of Philosophy", "Ph.D.", re.IGNORECASE),
    ]

    # Partition lines: search Education section first
    edu_header_re = re.compile(
        r"^(?:education|academic\s+background|academics|educational\s+qualifications?)\b", re.IGNORECASE
    )
    section_break_re = re.compile(
        r"^(?:experience|work\s+experience|projects|skills|technical\s+skills|certifications?|achievements?|activities)\b",
        re.IGNORECASE,
    )

    in_edu_section = False
    edu_lines: list[str] = []
    other_lines: list[str] = []

    for line in lines:
        if edu_header_re.search(line):
            in_edu_section = True
            continue
        if in_edu_section and section_break_re.search(line):
            in_edu_section = False
        if in_edu_section:
            edu_lines.append(line)
        else:
            other_lines.append(line)

    search_passes = [edu_lines, other_lines] if edu_lines else [lines]

    for line_set in search_passes:
        for line in line_set:
            if not degree:
                for pat, full_d, short_d, flags in degree_patterns:
                    if re.search(pat, line, flags):
                        degree = full_d
                        degree_short = short_d
                        if not evidence:
                            evidence = line
                        break

            if not institution:
                inst_match = re.search(
                    r"\b(?:(?:University|Institute|College|School)\s+of\s+[A-Za-z\s]+|"
                    r"[A-Z][a-zA-Z\s]+(?:Institute\s+of\s+(?:Technology|Science|Management)|College\s+of\s+[A-Za-z\s]+|University|Institute|College|Academy|IIT|NIT|IIIT|BITS|MIT|Stanford|Harvard))"
                    r"(?:\s*,\s*[A-Z][a-zA-Z\s]+)?\b",
                    line,
                )
                if inst_match:
                    institution = inst_match.group(0).strip()
                    if not evidence:
                        evidence = line

            if not grad_year:
                years = [int(y) for y in re.findall(r"\b(20[1-3][0-9])\b", line)]
                if years:
                    valid_years = [y for y in years if 2010 <= y <= 2035]
                    if valid_years:
                        grad_year = max(valid_years)
                        if not evidence:
                            evidence = line

            if not gpa:
                gpa_match = re.search(
                    r"(?:cgpa|gpa|cpi)\s*[:=-]?\s*([0-9]+(?:\.[0-9]+)?(?:\s*/\s*[0-9]+)?)", line, re.IGNORECASE
                )
                if gpa_match:
                    gpa = gpa_match.group(1).strip()

        if degree and institution and grad_year:
            break

    edu_dict = {
        "degree": degree,
        "degree_short": degree_short,
        "institution": institution,
        "graduation_year": grad_year,
        "gpa": gpa,
        "evidence": evidence,
        "verified": False,
    }
    return edu_dict, degree_short, evidence


# ---------------------------------------------------------------------------
# Experience Parsing Helpers
# ---------------------------------------------------------------------------

MONTH_MAP: dict[str, int] = {
    "jan": 1, "january": 1,
    "feb": 2, "february": 2,
    "mar": 3, "march": 3,
    "apr": 4, "april": 4,
    "may": 5,
    "jun": 6, "june": 6,
    "jul": 7, "july": 7,
    "aug": 8, "august": 8,
    "sep": 9, "september": 9,
    "oct": 10, "october": 10,
    "nov": 11, "november": 11,
    "dec": 12, "december": 12,
}

DATE_RANGE_RE = re.compile(
    r"\b((?:(?:jan|feb|mar|apr|may|jun|jul|aug|sep|oct|nov|dec)[a-z]*\.?\s*(?:20[1-3][0-9]|19[89][0-9])|"
    r"(?:0?[1-9]|1[0-2])/(?:20[1-3][0-9]|19[89][0-9])|"
    r"(?:20[1-3][0-9]|19[89][0-9])))\s*"
    r"[-–—to]+\s*"
    r"((?:(?:jan|feb|mar|apr|may|jun|jul|aug|sep|oct|nov|dec)[a-z]*\.?\s*(?:20[1-3][0-9]|19[89][0-9])|"
    r"(?:0?[1-9]|1[0-2])/(?:20[1-3][0-9]|19[89][0-9])|"
    r"(?:20[1-3][0-9]|19[89][0-9])|"
    r"present|current|now|ongoing))\b",
    re.IGNORECASE,
)


def _parse_endpoint(date_str: str, is_end: bool, today: datetime.date) -> tuple[int, int]:
    """Parse a date endpoint string to (year, month)."""
    clean = date_str.strip().lower()
    if clean in ("present", "current", "now", "ongoing"):
        return today.year, today.month

    # Check MM/YYYY
    m = re.match(r"^(\d{1,2})/(20[1-3][0-9]|19[89][0-9])$", clean)
    if m:
        return int(m.group(2)), int(m.group(1))

    # Check Mon YYYY
    m = re.match(r"^([a-z]+)\.?\s*(20[1-3][0-9]|19[89][0-9])$", clean)
    if m:
        month_name = m.group(1)
        year = int(m.group(2))
        month_num = MONTH_MAP.get(month_name[:3], 1 if not is_end else 12)
        return year, month_num

    # Check YYYY
    m = re.match(r"^(20[1-3][0-9]|19[89][0-9])$", clean)
    if m:
        year = int(m.group(1))
        return year, (12 if is_end else 1)

    return today.year, today.month


def _merge_intervals(intervals: list[tuple[int, int]]) -> list[tuple[int, int]]:
    """Merge overlapping or contiguous month intervals."""
    if not intervals:
        return []
    sorted_intervals = sorted(intervals, key=lambda x: (x[0], x[1]))
    merged = [list(sorted_intervals[0])]
    for start, end in sorted_intervals[1:]:
        if start <= merged[-1][1] + 1:
            merged[-1][1] = max(merged[-1][1], end)
        else:
            merged.append([start, end])
    return [(s, e) for s, e in merged]


def _parse_role_and_org(line: str, matched_role: Optional[str]) -> tuple[str, str]:
    """Parse role and organization from a line, separating company from role."""
    clean = re.sub(r"^(?:\d+[\.\)]\s*|[•\-\*]\s*)", "", line).strip()
    clean_no_date = DATE_RANGE_RE.sub("", clean).strip().rstrip("|-–—:, ")

    # "Role at Company"
    if re.search(r"\s+at\s+", clean_no_date, re.IGNORECASE):
        parts = re.split(r"\s+at\s+", clean_no_date, flags=re.IGNORECASE)
        part0 = parts[0].strip()
        part1 = " at ".join(parts[1:]).strip()
        if matched_role and matched_role.lower() in part0.lower():
            return matched_role.title(), part1
        elif matched_role and matched_role.lower() in part1.lower():
            return matched_role.title(), part0
        return part0.title(), part1

    # "Role | Company" or "Company | Role"
    if any(sep in clean_no_date for sep in ("|", " - ", " – ", " — ")):
        parts = [p.strip() for p in re.split(r"\s*[\|\–\—]|\s+\-\s+", clean_no_date) if p.strip()]
        if len(parts) >= 2:
            if matched_role:
                for idx, p in enumerate(parts):
                    if matched_role.lower() in p.lower():
                        org_parts = [parts[i] for i in range(len(parts)) if i != idx]
                        return matched_role.title(), " | ".join(org_parts)
            return parts[0].title(), parts[1]

    r = matched_role.title() if matched_role else "Technical Contributor"
    return r, clean_no_date


def extract_experience_from_text(
    text: str,
    today: Optional[datetime.date] = None,
) -> tuple[list[dict[str, Any]], float, str]:
    """Extract work experience entries, internships, duration, and estimate years of experience.

    Supports injectable reference date (`today`), unstructured resumes without an
    explicit Experience header, multi-line job entries, company/organization separation,
    month-level precision, and overlapping interval merging.
    """
    ref_date = today or datetime.date.today()
    experiences: list[dict[str, Any]] = []
    lines = [l.strip() for l in text.split("\n") if l.strip()]

    role_patterns = re.compile(
        r"\b(?:software\s+engineer\s+intern|software\s+developer\s+intern|engineering\s+intern|"
        r"data\s+science\s+intern|research\s+intern|frontend\s+developer\s+intern|backend\s+developer\s+intern|"
        r"senior\s+software\s+engineer|lead\s+software\s+engineer|software\s+engineer|"
        r"software\s+developer|backend\s+developer|frontend\s+developer|"
        r"full\s+stack\s+developer|devops\s+engineer|research\s+assistant|teaching\s+assistant|"
        r"intern(?:ship)?|sde\s*[-–—]?\s*(?:intern|[123]|i+)?|swe\s*[-–—]?\s*(?:intern|[123]|i+)?)\b",
        re.IGNORECASE,
    )

    exp_header_re = re.compile(
        r"^(?:experience|work\s+experience|professional\s+experience|internships?|employment\s+history)\b",
        re.IGNORECASE,
    )
    section_break_re = re.compile(
        r"^(?:education|projects|skills|technical\s+skills|certifications?|achievements?|activities)\b",
        re.IGNORECASE,
    )

    has_exp_section = any(exp_header_re.search(l) for l in lines)
    in_exp_section = not has_exp_section  # If no section header, scan candidate lines across doc

    current_entry: Optional[dict[str, Any]] = None

    for line in lines:
        if has_exp_section and exp_header_re.search(line):
            in_exp_section = True
            continue
        if has_exp_section and in_exp_section and section_break_re.search(line):
            in_exp_section = False
            if current_entry:
                experiences.append(current_entry)
                current_entry = None
            continue

        if in_exp_section:
            # If in unstructured mode, skip obvious education/skills/header lines
            if not has_exp_section and section_break_re.search(line):
                if current_entry:
                    experiences.append(current_entry)
                    current_entry = None
                continue

            role_m = role_patterns.search(line)
            date_m = DATE_RANGE_RE.search(line)

            if role_m or date_m:
                # Decide whether to attach to current_entry or start a new entry
                if current_entry:
                    has_role = current_entry["role"] != "Technical Contributor"
                    has_dur = bool(current_entry["duration"])

                    # If current entry already has role AND duration, this line is a new entry
                    if has_role and has_dur and (role_m or (date_m and not line.startswith(("-", "•", "*")))):
                        experiences.append(current_entry)
                        current_entry = None
                    elif not has_dur and date_m:
                        # Attach date to existing entry
                        current_entry["duration"] = date_m.group(0)
                        current_entry["evidence"] += " | " + line
                        if role_m and not has_role:
                            r_str, o_str = _parse_role_and_org(line, role_m.group(0))
                            current_entry["role"] = r_str
                            if o_str:
                                current_entry["organization"] = o_str
                        continue
                    elif not has_role and role_m:
                        r_str, o_str = _parse_role_and_org(line, role_m.group(0))
                        current_entry["role"] = r_str
                        if o_str:
                            current_entry["organization"] = o_str
                        current_entry["evidence"] += " | " + line
                        continue
                    else:
                        experiences.append(current_entry)
                        current_entry = None

                matched_role_str = role_m.group(0) if role_m else None
                r_str, o_str = _parse_role_and_org(line, matched_role_str)
                dur_str = date_m.group(0) if date_m else ""
                current_entry = {
                    "role": r_str,
                    "organization": o_str,
                    "duration": dur_str,
                    "description": "",
                    "evidence": line,
                    "verified": False,
                }
            elif current_entry:
                desc = line.lstrip("-•* ").strip()
                if not current_entry["description"]:
                    current_entry["description"] = desc
                else:
                    current_entry["description"] += " " + desc

    if current_entry:
        experiences.append(current_entry)

    # Calculate experience intervals with month-level precision
    intervals: list[tuple[int, int]] = []
    fallback_years = 0.0

    for exp in experiences:
        dur = exp.get("duration", "")
        dm = DATE_RANGE_RE.search(dur) if dur else None
        if dm:
            sy, sm = _parse_endpoint(dm.group(1), is_end=False, today=ref_date)
            ey, em = _parse_endpoint(dm.group(2), is_end=True, today=ref_date)
            s_idx = sy * 12 + (sm - 1)
            e_idx = ey * 12 + (em - 1)
            if e_idx >= s_idx:
                intervals.append((s_idx, e_idx))
            else:
                intervals.append((s_idx, s_idx))
        else:
            # Undated fallback: internships get 0.25, others 0.5
            is_intern = "intern" in exp.get("role", "").lower() or "intern" in exp.get("organization", "").lower()
            exp["fallback_dates"] = True
            fallback_years += 0.25 if is_intern else 0.5

    merged = _merge_intervals(intervals)
    total_dated_months = sum(e - s + 1 for s, e in merged)
    est_years = round((total_dated_months / 12.0) + fallback_years, 2)

    evidence_str = " | ".join(f"{e['role']} at {e['organization']}" for e in experiences[:2]) if experiences else ""
    return experiences, est_years, evidence_str


def extract_certifications_achievements_from_text(text: str) -> tuple[list[dict[str, Any]], str]:
    """Extract explicitly stated certifications, hackathon placements, and honors."""
    items: list[dict[str, Any]] = []
    lines = [l.strip() for l in text.split("\n") if l.strip()]

    cert_patterns: list[tuple[str, str]] = [
        (r"\baws\s+certified\b[^\n,]*", "certification"),
        (r"\bgoogle\s+cloud\s+certified\b[^\n,]*", "certification"),
        (r"\bmicrosoft\s+certified\b[^\n,]*", "certification"),
        (r"\bcertified\s+kubernetes\b[^\n,]*", "certification"),
        (
            r"(?:(?:winner|runner-?up|finalist|champion|1st\s+place|first\s+place|2nd\s+place|second\s+place|3rd\s+place|third\s+place)"
            r"[^\n\r]*?hackathon[^\n\r]*|"
            r"[^\n\r]*?hackathon[^\n\r]*?(?:winner|runner-?up|finalist|champion|1st\s+place|first\s+place|2nd\s+place|second\s+place|3rd\s+place|third\s+place)[^\n\r]*)",
            "achievement",
        ),
        (r"\bdean'?s\s+list\b", "achievement"),
    ]

    for line in lines:
        clean_line = re.sub(r"^(?:\d+[\.\)]\s*|[•\-\*]\s*)", "", line).strip()
        for pat, item_type in cert_patterns:
            m = re.search(pat, clean_line, re.IGNORECASE)
            if m:
                raw_title = m.group(0).strip().strip("-|,:() ")
                title = re.sub(r"^(?:\d+[\.\)]\s*|[•\-\*]\s*)", "", raw_title).strip()
                if not any(i["title"].lower() == title.lower() for i in items):
                    items.append({
                        "title": title,
                        "type": item_type,
                        "evidence": line,
                        "verified": False,
                    })

    evidence_summary = " | ".join(i["title"] for i in items[:3]) if items else ""
    return items, evidence_summary


def extract_external_links_from_text(text: str) -> tuple[dict[str, list[str]], str]:
    """Extract GitHub, portfolio, LinkedIn, and competitive programming links from resume text."""
    github_urls: list[str] = []
    linkedin_urls: list[str] = []
    competitive_programming_urls: list[str] = []
    portfolio_urls: list[str] = []
    other_urls: list[str] = []

    # Match URLs with protocol or schemeless domains
    url_re = re.compile(
        r"(?:https?://|www\.)[A-Za-z0-9_\-\.\/\?\=\&\#\%]+|"
        r"\b(?:github\.com|linkedin\.com|(?:leetcode|codeforces|codechef|hackerrank|kaggle)\.com|atcoder\.jp)"
        r"/[A-Za-z0-9_\-\.\/\?\=\&\#\%]+|"
        r"\b[A-Za-z0-9_\-]+\.(?:dev|me|io|app|tech|page)(?:/[A-Za-z0-9_\-\.\/\?\=\&\#\%]*)?",
        re.IGNORECASE,
    )

    raw_matches = url_re.findall(text)
    for raw in raw_matches:
        url = raw.strip().rstrip(".,;:!?)>\"'")
        if not url:
            continue
        if not url.startswith(("http://", "https://")):
            url = "https://" + url

        url_lower = url.lower()
        if "github.com" in url_lower or "github.io" in url_lower:
            if url not in github_urls:
                github_urls.append(url)
        elif "linkedin.com" in url_lower:
            if url not in linkedin_urls:
                linkedin_urls.append(url)
        elif any(cp in url_lower for cp in ["leetcode.com", "codeforces.com", "codechef.com", "hackerrank.com", "atcoder.jp", "kaggle.com"]):
            if url not in competitive_programming_urls:
                competitive_programming_urls.append(url)
        elif any(ext in url_lower for ext in [".dev", ".me", ".io", ".app", ".tech", ".page", "portfolio", "personal"]):
            if url not in portfolio_urls:
                portfolio_urls.append(url)
        else:
            if url not in other_urls:
                other_urls.append(url)

    link_data = {
        "github_urls": github_urls,
        "portfolio_urls": portfolio_urls,
        "linkedin_urls": linkedin_urls,
        "competitive_programming_urls": competitive_programming_urls,
        "other_urls": other_urls,
        "note": "Extracted from resume document; not independently verified.",
    }
    evidence_summary = (
        f"GitHub: {len(github_urls)}, LinkedIn: {len(linkedin_urls)}, "
        f"Portfolio: {len(portfolio_urls)}, CP: {len(competitive_programming_urls)}"
    )
    return link_data, evidence_summary


# ---------------------------------------------------------------------------
# Pluggable ResumeSignalExtractor Class
# ---------------------------------------------------------------------------


class ResumeSignalExtractor(BaseSignalExtractor):
    """Signal extractor for parsing candidate PDF and DOCX resumes."""

    @property
    def descriptor(self) -> ExtractorDescriptor:
        return ExtractorDescriptor(
            extractor_id="resume_parser",
            name="Resume & CV Signal Analyzer",
            description="Extracts technical skills, projects, education, experience, achievements, and links from PDF and DOCX resumes.",
            category="experience",
            required_inputs=["resume_url"],
            optional_inputs=["resume_text", "resume_bytes", "custom_fields"],
            produced_signals=[
                "resume_extracted_skills",
                "resume_programming_languages",
                "resume_frameworks_libraries",
                "resume_databases",
                "resume_tools_technologies",
                "resume_projects",
                "resume_education",
                "resume_degree_level",
                "resume_experience",
                "resume_experience_years",
                "resume_certifications_achievements",
                "resume_external_links",
            ],
            signals=[
                SignalDefinition(
                    signal_key="resume_extracted_skills",
                    name="Extracted Skills",
                    description="Categorized technical skills identified in resume text (languages, frameworks, databases, tools).",
                    category="experience",
                    value_type="dict",
                ),
                SignalDefinition(
                    signal_key="resume_programming_languages",
                    name="Programming Languages",
                    description="Programming languages explicitly identified in resume text.",
                    category="experience",
                    value_type="list",
                ),
                SignalDefinition(
                    signal_key="resume_frameworks_libraries",
                    name="Frameworks & Libraries",
                    description="Software frameworks and libraries mentioned in resume text.",
                    category="experience",
                    value_type="list",
                ),
                SignalDefinition(
                    signal_key="resume_databases",
                    name="Databases & Data Stores",
                    description="Databases and storage systems explicitly listed in resume text.",
                    category="experience",
                    value_type="list",
                ),
                SignalDefinition(
                    signal_key="resume_tools_technologies",
                    name="Development Tools & Platforms",
                    description="Tools, cloud platforms, and DevOps technologies found in resume text.",
                    category="experience",
                    value_type="list",
                ),
                SignalDefinition(
                    signal_key="resume_projects",
                    name="Projects",
                    description="Candidate projects with titles, descriptions, technologies, and evidence snippets.",
                    category="experience",
                    value_type="list",
                ),
                SignalDefinition(
                    signal_key="resume_education",
                    name="Education Background",
                    description="Degrees, educational institutions, graduation years, and academic standing.",
                    category="experience",
                    value_type="dict",
                ),
                SignalDefinition(
                    signal_key="resume_degree_level",
                    name="Highest Degree Level",
                    description="Highest academic degree level identified (e.g. B.Tech, M.Tech, B.S.).",
                    category="experience",
                    value_type="string",
                ),
                SignalDefinition(
                    signal_key="resume_experience",
                    name="Work & Internship Experience",
                    description="Internships and professional roles with organizations, titles, durations, and details.",
                    category="experience",
                    value_type="list",
                ),
                SignalDefinition(
                    signal_key="resume_experience_years",
                    name="Years of Experience",
                    description="Estimated duration or years of professional and internship experience.",
                    category="experience",
                    value_type="float",
                    unit="years",
                ),
                SignalDefinition(
                    signal_key="resume_certifications_achievements",
                    name="Certifications & Achievements",
                    description="Certifications, hackathon placements, awards, and honors.",
                    category="experience",
                    value_type="list",
                ),
                SignalDefinition(
                    signal_key="resume_external_links",
                    name="External Professional Links",
                    description="Links to GitHub, LinkedIn, portfolio, and competitive programming profiles.",
                    category="experience",
                    value_type="dict",
                ),
            ],
            is_enabled_by_default=False,
            optional_config={
                "max_resume_bytes": "Maximum resume file size in bytes (default 10 MB)",
                "download_timeout": "Download timeout in seconds (default 10.0)",
                "allow_local_files": "Allow reading local file paths for testing (default False)",
                "resume_url_field": "Custom source column name for resume URL in registration data",
            },
        )

    def extract(
        self,
        member: MemberSchema,
        selected_signals: Optional[set[str]] = None,
        config: Optional[dict[str, Any]] = None,
    ) -> list[ExtractedSignal]:
        """Produce signals for the given member."""
        target_signals = selected_signals or set(self.descriptor.produced_signals)
        cfg = config or {}
        results: list[ExtractedSignal] = []
        now_ts = datetime.datetime.now(datetime.timezone.utc).isoformat()

        # Resolve resume input
        resume_val, source_field = self._resolve_resume_input(member, cfg)

        # Handle missing input gracefully
        if not resume_val:
            for s_key in target_signals:
                results.append(
                    ExtractedSignal(
                        signal_key=s_key,
                        category="experience",
                        status=SignalStatus.NOT_APPLICABLE,
                        value=None,
                        error_message="Member did not provide a resume URL or document.",
                    )
                )
            return results

        # Load document bytes or pre-extracted text
        raw_bytes: Optional[bytes] = None
        extracted_text: Optional[str] = None
        load_error: Optional[str] = None
        doc_metadata: dict[str, Any] = {}

        # 1. Check if raw text is already provided in custom fields (e.g. test fixtures)
        if isinstance(resume_val, str) and member.custom_fields.get("resume_text"):
            extracted_text = member.custom_fields["resume_text"]
            doc_metadata["format"] = "text_override"

        # 2. Check if raw bytes are directly passed in custom fields
        elif member.custom_fields.get("resume_bytes"):
            raw_bytes = member.custom_fields["resume_bytes"]

        # 3. Handle URL or local file path
        elif isinstance(resume_val, str):
            res_str = resume_val.strip()
            if res_str.startswith(("http://", "https://")):
                max_bytes = int(cfg.get("max_resume_bytes", DEFAULT_MAX_RESUME_BYTES))
                timeout = float(cfg.get("download_timeout", DEFAULT_DOWNLOAD_TIMEOUT))
                raw_bytes, load_error = download_resume_file(res_str, max_bytes=max_bytes, timeout=timeout)
            elif not cfg.get("allow_local_files", False):
                # When allow_local_files is False (default):
                # Never call os.path.exists() on untrusted inputs (prevents UNC SMB hangs & leaks).
                load_error = "Local file access is disabled."
            else:
                # allow_local_files is explicitly True
                if _is_unc_path(res_str):
                    load_error = "Local file access is disabled."
                else:
                    if res_str.lower().startswith("file://"):
                        parsed = urllib.parse.urlparse(res_str)
                        local_path = urllib.parse.unquote(parsed.path)
                        if os.name == "nt" and re.match(r"^/[a-zA-Z]:", local_path):
                            local_path = local_path.lstrip("/")
                    else:
                        local_path = res_str

                    if os.path.exists(local_path):
                        try:
                            with open(local_path, "rb") as f:
                                raw_bytes = f.read()
                        except Exception as exc:
                            load_error = f"Failed to read local resume file: {exc}"
                    else:
                        load_error = f"Local resume file not found: {local_path}"

        # Handle loading error
        if load_error or (raw_bytes is None and extracted_text is None):
            err = load_error or "Unable to acquire resume document."
            for s_key in target_signals:
                results.append(
                    ExtractedSignal(
                        signal_key=s_key,
                        category="experience",
                        status=SignalStatus.FAILED,
                        value=None,
                        error_message=err,
                        evidence=SignalEvidence(
                            source="resume_parser",
                            raw_input=str(resume_val),
                            extracted_at=now_ts,
                            metadata={"error": err},
                        ),
                    )
                )
            return results

        # Extract text if not already available
        if extracted_text is None and raw_bytes is not None:
            filename = str(resume_val) if isinstance(resume_val, str) else "resume.pdf"
            extracted_text, parse_error, doc_metadata = extract_text_from_document(raw_bytes, filename)
            if parse_error or not extracted_text:
                err = parse_error or "Failed to extract text from resume."
                for s_key in target_signals:
                    results.append(
                        ExtractedSignal(
                            signal_key=s_key,
                            category="experience",
                            status=SignalStatus.FAILED,
                            value=None,
                            error_message=err,
                            evidence=SignalEvidence(
                                source="resume_parser",
                                raw_input=str(resume_val),
                                extracted_at=now_ts,
                                metadata=doc_metadata,
                            ),
                        )
                    )
                return results

        # Extract all component signals from document text
        skills_dict, skill_evidence = extract_skills_from_text(extracted_text)
        projects_list, proj_evidence = extract_projects_from_text(extracted_text)
        edu_dict, degree_level, edu_evidence = extract_education_from_text(extracted_text)
        exp_list, exp_years, exp_evidence = extract_experience_from_text(extracted_text)
        achieve_list, achieve_evidence = extract_certifications_achievements_from_text(extracted_text)
        links_dict, links_evidence = extract_external_links_from_text(extracted_text)

        # Dynamic confidence heuristics based on extraction quality and signal evidence
        char_count = len(extracted_text) if extracted_text else 0
        doc_factor = 1.0 if char_count >= 500 else (0.85 if char_count >= 200 else 0.70)

        # Build individual ExtractedSignal responses based on target_signals
        if "resume_extracted_skills" in target_signals:
            cats = [skills_dict["languages"], skills_dict["frameworks"], skills_dict["databases"], skills_dict["tools"]]
            cat_count = sum(1 for c in cats if c)
            if cat_count >= 3:
                s_conf = round(min(1.0, 0.90 * doc_factor), 2)
                s_rule = "multicategory_skills_identified"
            elif cat_count >= 1:
                s_conf = round(min(1.0, 0.80 * doc_factor), 2)
                s_rule = "partial_category_skills_identified"
            else:
                s_conf = 0.50
                s_rule = "no_skills_detected"

            results.append(
                ExtractedSignal(
                    signal_key="resume_extracted_skills",
                    category="experience",
                    status=SignalStatus.EXTRACTED,
                    value=skills_dict,
                    evidence=SignalEvidence(
                        source="resume_parser",
                        raw_input=str(resume_val),
                        extracted_at=now_ts,
                        confidence=s_conf,
                        metadata={
                            "confidence_rule": s_rule,
                            "total_skills": skills_dict["total_skills_count"],
                            "snippets": skill_evidence,
                            **doc_metadata,
                        },
                    ),
                )
            )

        if "resume_programming_languages" in target_signals:
            lang_count = len(skills_dict["languages"])
            if lang_count >= 2:
                l_conf = round(min(1.0, 0.90 * doc_factor), 2)
                l_rule = "multiple_languages_found"
            elif lang_count == 1:
                l_conf = round(min(1.0, 0.80 * doc_factor), 2)
                l_rule = "single_language_found"
            else:
                l_conf = 0.50
                l_rule = "no_languages_found"

            results.append(
                ExtractedSignal(
                    signal_key="resume_programming_languages",
                    category="experience",
                    status=SignalStatus.EXTRACTED,
                    value=skills_dict["languages"],
                    evidence=SignalEvidence(
                        source="resume_parser",
                        raw_input=str(resume_val),
                        extracted_at=now_ts,
                        confidence=l_conf,
                        metadata={"confidence_rule": l_rule, "count": lang_count},
                    ),
                )
            )

        if "resume_frameworks_libraries" in target_signals:
            fw_count = len(skills_dict["frameworks"])
            if fw_count >= 2:
                f_conf = round(min(1.0, 0.90 * doc_factor), 2)
                f_rule = "multiple_frameworks_found"
            elif fw_count == 1:
                f_conf = round(min(1.0, 0.80 * doc_factor), 2)
                f_rule = "single_framework_found"
            else:
                f_conf = 0.50
                f_rule = "no_frameworks_found"

            results.append(
                ExtractedSignal(
                    signal_key="resume_frameworks_libraries",
                    category="experience",
                    status=SignalStatus.EXTRACTED,
                    value=skills_dict["frameworks"],
                    evidence=SignalEvidence(
                        source="resume_parser",
                        raw_input=str(resume_val),
                        extracted_at=now_ts,
                        confidence=f_conf,
                        metadata={"confidence_rule": f_rule, "count": fw_count},
                    ),
                )
            )

        if "resume_databases" in target_signals:
            db_count = len(skills_dict["databases"])
            if db_count >= 1:
                d_conf = round(min(1.0, 0.85 * doc_factor), 2)
                d_rule = "databases_found"
            else:
                d_conf = 0.50
                d_rule = "no_databases_found"

            results.append(
                ExtractedSignal(
                    signal_key="resume_databases",
                    category="experience",
                    status=SignalStatus.EXTRACTED,
                    value=skills_dict["databases"],
                    evidence=SignalEvidence(
                        source="resume_parser",
                        raw_input=str(resume_val),
                        extracted_at=now_ts,
                        confidence=d_conf,
                        metadata={"confidence_rule": d_rule, "count": db_count},
                    ),
                )
            )

        if "resume_tools_technologies" in target_signals:
            t_count = len(skills_dict["tools"])
            if t_count >= 1:
                t_conf = round(min(1.0, 0.85 * doc_factor), 2)
                t_rule = "tools_found"
            else:
                t_conf = 0.50
                t_rule = "no_tools_found"

            results.append(
                ExtractedSignal(
                    signal_key="resume_tools_technologies",
                    category="experience",
                    status=SignalStatus.EXTRACTED,
                    value=skills_dict["tools"],
                    evidence=SignalEvidence(
                        source="resume_parser",
                        raw_input=str(resume_val),
                        extracted_at=now_ts,
                        confidence=t_conf,
                        metadata={"confidence_rule": t_rule, "count": t_count},
                    ),
                )
            )

        if "resume_projects" in target_signals:
            p_count = len(projects_list)
            if any(p.get("technologies") and p.get("description") for p in projects_list):
                p_conf = round(min(1.0, 0.85 * doc_factor), 2)
                p_rule = "detailed_projects_with_tech"
            elif p_count > 0:
                p_conf = round(min(1.0, 0.75 * doc_factor), 2)
                p_rule = "projects_found"
            else:
                p_conf = 0.50
                p_rule = "no_projects_found"

            results.append(
                ExtractedSignal(
                    signal_key="resume_projects",
                    category="experience",
                    status=SignalStatus.EXTRACTED,
                    value=projects_list,
                    evidence=SignalEvidence(
                        source="resume_parser",
                        raw_input=str(resume_val),
                        extracted_at=now_ts,
                        confidence=p_conf,
                        metadata={"confidence_rule": p_rule, "projects_count": p_count, "sample": proj_evidence},
                    ),
                )
            )

        if "resume_education" in target_signals:
            has_deg = bool(edu_dict.get("degree"))
            has_inst = bool(edu_dict.get("institution"))
            if has_deg and has_inst:
                e_conf = round(min(1.0, 0.90 * doc_factor), 2)
                e_rule = "degree_and_institution_found"
            elif has_deg:
                e_conf = round(min(1.0, 0.70 * doc_factor), 2)
                e_rule = "degree_only_found"
            elif has_inst:
                e_conf = round(min(1.0, 0.60 * doc_factor), 2)
                e_rule = "institution_only_found"
            else:
                e_conf = 0.40
                e_rule = "no_education_found"

            results.append(
                ExtractedSignal(
                    signal_key="resume_education",
                    category="experience",
                    status=SignalStatus.EXTRACTED,
                    value=edu_dict,
                    evidence=SignalEvidence(
                        source="resume_parser",
                        raw_input=str(resume_val),
                        extracted_at=now_ts,
                        confidence=e_conf,
                        metadata={"confidence_rule": e_rule, "evidence": edu_evidence},
                    ),
                )
            )

        if "resume_degree_level" in target_signals:
            if degree_level:
                deg_conf = round(min(1.0, 0.90 * doc_factor), 2)
                deg_rule = "degree_level_identified"
            else:
                deg_conf = 0.40
                deg_rule = "degree_level_unidentified"

            results.append(
                ExtractedSignal(
                    signal_key="resume_degree_level",
                    category="experience",
                    status=SignalStatus.EXTRACTED,
                    value=degree_level,
                    evidence=SignalEvidence(
                        source="resume_parser",
                        raw_input=str(resume_val),
                        extracted_at=now_ts,
                        confidence=deg_conf,
                        metadata={"confidence_rule": deg_rule, "degree_identified": degree_level},
                    ),
                )
            )

        if "resume_experience" in target_signals:
            has_dated = any(e.get("duration") and not e.get("fallback_dates") for e in exp_list)
            if has_dated:
                exp_conf = round(min(1.0, 0.85 * doc_factor), 2)
                exp_rule = "dated_experience_entries"
            elif exp_list:
                exp_conf = round(min(1.0, 0.50 * doc_factor), 2)
                exp_rule = "undated_experience_fallback"
            else:
                exp_conf = 0.40
                exp_rule = "no_experience_entries"

            results.append(
                ExtractedSignal(
                    signal_key="resume_experience",
                    category="experience",
                    status=SignalStatus.EXTRACTED,
                    value=exp_list,
                    evidence=SignalEvidence(
                        source="resume_parser",
                        raw_input=str(resume_val),
                        extracted_at=now_ts,
                        confidence=exp_conf,
                        metadata={"confidence_rule": exp_rule, "entries_count": len(exp_list), "sample": exp_evidence},
                    ),
                )
            )

        if "resume_experience_years" in target_signals:
            has_dated = any(e.get("duration") and not e.get("fallback_dates") for e in exp_list)
            if exp_years > 0 and has_dated:
                ey_conf = round(min(1.0, 0.85 * doc_factor), 2)
                ey_rule = "calculated_from_dated_intervals"
            elif exp_years > 0:
                ey_conf = round(min(1.0, 0.45 * doc_factor), 2)
                ey_rule = "estimated_from_fallback"
            else:
                ey_conf = 0.50
                ey_rule = "zero_experience_years"

            results.append(
                ExtractedSignal(
                    signal_key="resume_experience_years",
                    category="experience",
                    status=SignalStatus.EXTRACTED,
                    value=exp_years,
                    evidence=SignalEvidence(
                        source="resume_parser",
                        raw_input=str(resume_val),
                        extracted_at=now_ts,
                        confidence=ey_conf,
                        metadata={"confidence_rule": ey_rule, "estimated_years": exp_years},
                    ),
                )
            )

        if "resume_certifications_achievements" in target_signals:
            a_count = len(achieve_list)
            if a_count > 0:
                ac_conf = round(min(1.0, 0.85 * doc_factor), 2)
                ac_rule = "verified_achievement_patterns"
            else:
                ac_conf = 0.50
                ac_rule = "no_achievements_found"

            results.append(
                ExtractedSignal(
                    signal_key="resume_certifications_achievements",
                    category="experience",
                    status=SignalStatus.EXTRACTED,
                    value=achieve_list,
                    evidence=SignalEvidence(
                        source="resume_parser",
                        raw_input=str(resume_val),
                        extracted_at=now_ts,
                        confidence=ac_conf,
                        metadata={"confidence_rule": ac_rule, "count": a_count, "sample": achieve_evidence},
                    ),
                )
            )

        if "resume_external_links" in target_signals:
            total_links = (
                len(links_dict.get("github_urls", []))
                + len(links_dict.get("linkedin_urls", []))
                + len(links_dict.get("portfolio_urls", []))
                + len(links_dict.get("competitive_programming_urls", []))
                + len(links_dict.get("other_urls", []))
            )
            if total_links > 0:
                lk_conf = round(min(1.0, 0.95 * doc_factor), 2)
                lk_rule = "validated_urls_with_recognized_domains"
            else:
                lk_conf = 0.50
                lk_rule = "no_links_found"

            results.append(
                ExtractedSignal(
                    signal_key="resume_external_links",
                    category="experience",
                    status=SignalStatus.EXTRACTED,
                    value=links_dict,
                    evidence=SignalEvidence(
                        source="resume_parser",
                        raw_input=str(resume_val),
                        extracted_at=now_ts,
                        confidence=lk_conf,
                        metadata={"confidence_rule": lk_rule, "summary": links_evidence},
                    ),
                )
            )

        return results

    def _resolve_resume_input(
        self,
        member: MemberSchema,
        config: dict[str, Any],
    ) -> tuple[Optional[Any], str]:
        """Resolve the raw resume URL, document reference, or text from member data."""
        custom_field_key = config.get("resume_url_field")
        if custom_field_key and custom_field_key in member.custom_fields:
            return member.custom_fields[custom_field_key], custom_field_key

        if member.resume_url:
            return member.resume_url, "resume_url"

        for alt_key in ["resume_url", "resume", "cv", "resume_link", "resume_file"]:
            if alt_key in member.custom_fields and member.custom_fields[alt_key]:
                return member.custom_fields[alt_key], alt_key

        if member.custom_fields.get("resume_bytes"):
            return member.custom_fields["resume_bytes"], "resume_bytes"

        if member.custom_fields.get("resume_text"):
            return member.custom_fields["resume_text"], "resume_text"

        return None, "resume_url"
